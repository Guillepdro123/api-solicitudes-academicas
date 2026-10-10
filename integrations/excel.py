import io
import json
import logging
import os
import threading

import openpyxl
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload

import metricas

logger = logging.getLogger(__name__)
SCOPES = ["https://www.googleapis.com/auth/drive"]
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Hoja del libro donde se guardan las solicitudes (se crea si no existe)
HOJA = "Solicitudes"
COLUMNAS = ["numero", "id", "estudiante", "tipo", "descripcion", "estado", "fecha"]
TIPOS = ["Certificado", "Constancia", "Homologación", "Cancelación de asignatura", "Reintegro", "Otro"]
ESTADOS = ["Pendiente", "En proceso", "Resuelta", "Rechazada"]

# El cliente de Google (httplib2) no es seguro entre hilos: se crea uno por hilo
_local = threading.local()


def _credenciales():
    """
    Lee la cuenta de servicio desde GOOGLE_CREDENTIALS_JSON (contenido del JSON, útil al
    desplegar) o desde el archivo indicado en GOOGLE_APPLICATION_CREDENTIALS.
    """
    contenido = os.getenv("GOOGLE_CREDENTIALS_JSON")
    if contenido:
        return service_account.Credentials.from_service_account_info(json.loads(contenido), scopes=SCOPES)

    ruta = os.getenv("GOOGLE_APPLICATION_CREDENTIALS", "credenciales.json")
    if not os.path.exists(ruta):
        raise FileNotFoundError(f"El archivo de credenciales {ruta} no existe.")
    return service_account.Credentials.from_service_account_file(ruta, scopes=SCOPES)


def _servicio():
    if not hasattr(_local, "servicio"):
        _local.servicio = build("drive", "v3", credentials=_credenciales(), cache_discovery=False)
    return _local.servicio


def _file_id():
    file_id = os.getenv("GOOGLE_DRIVE_FILE_ID")
    if not file_id or file_id == "TU_ID_DE_ARCHIVO_EN_DRIVE":
        raise ValueError("GOOGLE_DRIVE_FILE_ID no está configurado.")
    return file_id


def _descargar_libro():
    """Descarga el Excel de Drive a memoria, sin archivos temporales en disco."""
    request = _servicio().files().get_media(fileId=_file_id())
    stream = io.BytesIO()
    downloader = MediaIoBaseDownload(stream, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    stream.seek(0)
    return openpyxl.load_workbook(stream)


def _subir_libro(wb):
    """Sobrescribe el archivo de Drive con el libro modificado."""
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    media = MediaIoBaseUpload(stream, mimetype=XLSX_MIME, resumable=False)
    _servicio().files().update(fileId=_file_id(), media_body=media).execute()


def _hoja(wb, crear: bool):
    if HOJA in wb.sheetnames:
        return wb[HOJA]
    if not crear:
        return None
    ws = wb.create_sheet(HOJA)
    ws.append(COLUMNAS)
    return ws


def _fila_a_dict(valores):
    return {col: "" if v is None else str(v) for col, v in zip(COLUMNAS, valores)}


def leer_solicitudes() -> list[dict]:
    with metricas.excel.medir():
        return _leer_solicitudes()


def aplicar_operaciones(operaciones: list[dict]) -> list[dict | None]:
    """
    Descarga el libro una sola vez, aplica las operaciones en orden y lo sube una sola vez.

    Cada operación es {"tipo": "crear", "solicitud": {...}} o
    {"tipo": "actualizar", "id": "...", "estado": "..."}.
    Devuelve, por operación, la solicitud resultante (o None si el id no existe).
    """
    with metricas.excel.medir():
        return _aplicar_operaciones(operaciones)


def ultimo_numero() -> int:
    """Mayor número consecutivo guardado (0 si la hoja está vacía), para continuar la numeración."""
    numeros = [int(s["numero"]) for s in leer_solicitudes() if s["numero"].isdigit()]
    return max(numeros, default=0)


def libro_vacio():
    """Libro nuevo con la hoja «Solicitudes» vacía, encabezados con formato y listas desplegables."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = HOJA
    ws.append(COLUMNAS)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F3A5F")
        celda.alignment = Alignment(horizontal="center", vertical="center")
    for letra, ancho in zip("ABCDEFG", [10, 12, 30, 26, 40, 14, 26]):
        ws.column_dimensions[letra].width = ancho
    ws.freeze_panes = "A2"
    listas = {"tipo": TIPOS, "estado": ESTADOS}
    for columna, valores in listas.items():
        letra = "ABCDEFG"[COLUMNAS.index(columna)]
        dv = DataValidation(type="list", formula1='"' + ",".join(valores) + '"', allow_blank=True)
        dv.add(f"{letra}2:{letra}5000")
        ws.add_data_validation(dv)
    return wb


def reiniciar_drive():
    """Reemplaza el Excel de Drive por un libro vacío (mismo archivo, mismo enlace y permisos)."""
    with metricas.excel.medir():
        _subir_libro(libro_vacio())


def _leer_solicitudes() -> list[dict]:
    wb = _descargar_libro()
    ws = _hoja(wb, crear=False)
    if ws is None:
        return []
    i = COLUMNAS.index("id")
    return [_fila_a_dict(fila) for fila in ws.iter_rows(min_row=2, values_only=True) if fila and len(fila) > i and fila[i]]


def _aplicar_operaciones(operaciones: list[dict]) -> list[dict | None]:
    wb = _descargar_libro()
    ws = _hoja(wb, crear=True)
    col_id = COLUMNAS.index("id") + 1
    filas = {
        str(celda.value): celda.row
        for (celda,) in ws.iter_rows(min_row=2, min_col=col_id, max_col=col_id)
        if celda.value
    }
    col_estado = COLUMNAS.index("estado") + 1

    resultados = []
    hubo_cambios = False
    for op in operaciones:
        if op["tipo"] == "crear":
            solicitud = op["solicitud"]
            # Si un lote se reintenta después de haberse subido, no se duplica la fila
            if solicitud["id"] not in filas:
                ws.append([solicitud[col] for col in COLUMNAS])
                for celda in ws[ws.max_row]:
                    # Evita que un texto que empiece por "=" se guarde como fórmula
                    if isinstance(celda.value, str):
                        celda.data_type = "s"
                filas[solicitud["id"]] = ws.max_row
                hubo_cambios = True
            resultados.append(solicitud)
        else:
            fila = filas.get(op["id"])
            if fila is None:
                resultados.append(None)
                continue
            ws.cell(row=fila, column=col_estado, value=op["estado"])
            hubo_cambios = True
            resultados.append(_fila_a_dict(celda.value for celda in ws[fila][: len(COLUMNAS)]))

    if hubo_cambios:
        _subir_libro(wb)
        logger.info("Excel actualizado en Drive (%d operaciones).", len(operaciones))
    wb.close()
    return resultados
